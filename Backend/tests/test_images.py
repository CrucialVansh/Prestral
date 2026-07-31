"""Tests for embedded PPTX picture extraction and multimodal LLM wiring."""

from __future__ import annotations

import json
from io import BytesIO
from unittest.mock import MagicMock, patch

from app.config import Settings
from app.models.schemas import (
    BBox,
    Component,
    ComponentType,
    CreateSessionRequest,
    Deck,
    DocChunk,
    QueryMode,
    QueryRequest,
    SendMessageRequest,
    Slide,
)
from app.services.analysis import analyze_deck
from app.services.chat import get_or_create_session, send_message
from app.services.llm import LLMClient
from app.services.qa import answer_query
from app.services.slide_parser import parse_slides
from app.store import deck_store, session_store
from pptx import Presentation
from pptx.util import Inches


def _tiny_png_bytes() -> bytes:
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )


def _pptx_with_picture() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    box.text_frame.text = "Infographic slide"

    stream = BytesIO(_tiny_png_bytes())
    slide.shapes.add_picture(stream, Inches(1), Inches(1.5), width=Inches(4))

    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("The infographic shows Q3 revenue up 12% driven by enterprise.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extracts_picture_as_data_uri() -> None:
    parsed = parse_slides(_pptx_with_picture())
    pictures = [c for c in parsed.slides[0].components if c.type.value == "picture"]
    assert len(pictures) == 1
    pic = pictures[0]
    assert pic.has_image is True
    assert pic.id in parsed.images
    uri = parsed.images[pic.id]
    assert uri.startswith("data:image/")
    assert ";base64," in uri


def test_text_only_deck_has_no_images(sample_pptx_bytes: bytes) -> None:
    parsed = parse_slides(sample_pptx_bytes)
    assert parsed.images == {}
    assert all(not c.has_image for s in parsed.slides for c in s.components)


def test_response_schema_exposes_has_image_not_base64(client) -> None:
    schemas = client.get("/openapi.json").json()["components"]["schemas"]
    props = schemas["Component"]["properties"]
    assert "has_image" in props
    assert "image_data_uri" not in props
    assert "component_images" not in schemas.get("DeckAnalysisResponse", {}).get(
        "properties", {}
    )


def test_relate_slide_components_attaches_image_parts() -> None:
    """Analysis prompt for a picture component includes a multimodal image part."""
    settings = Settings(mistral_api_key="test-key")
    llm = LLMClient(settings)

    slide = Slide(
        index=0,
        components=[
            Component(
                id="slide0-shape1",
                type=ComponentType.PICTURE,
                bbox=BBox(left=0.1, top=0.1, width=0.4, height=0.4),
                text="Chart image",
                has_image=True,
            )
        ],
    )
    data_uri = "data:image/png;base64,abc123"
    chunks = [DocChunk(text="Revenue grew 12%.", source="paragraph:0")]

    captured: dict = {}

    def fake_chat(messages, *, json_mode=False, temperature=0.3):
        captured["messages"] = messages
        return json.dumps(
            {
                "slide0-shape1": {
                    "context": "Shows revenue growth.",
                    "sources": ["paragraph:0"],
                }
            }
        )

    with patch.object(llm, "chat", side_effect=fake_chat):
        result = llm.relate_slide_components(
            slide, chunks, images={"slide0-shape1": data_uri}
        )

    assert "slide0-shape1" in result
    user_content = captured["messages"][1]["content"]
    assert isinstance(user_content, list)
    image_parts = [p for p in user_content if p.get("type") == "image_url"]
    assert len(image_parts) == 1
    assert image_parts[0]["image_url"] == data_uri


def test_analyze_deck_stores_images_and_sets_has_image() -> None:
    """Full analyze path extracts pictures, stores them on Deck, marks has_image."""
    settings = Settings(mistral_api_key="test-key")

    class FakeEmbed:
        def __init__(self, _settings):
            pass

        def embed(self, texts):
            return [[1.0, 0.0] for _ in texts]

        def embed_chunks(self, chunks):
            for c in chunks:
                c.embedding = [1.0, 0.0]
            return chunks

    class FakeLLM:
        def __init__(self, _settings):
            pass

        def relate_slide_components(self, slide, retrieved, images=None):
            # Prove images were passed for picture components
            assert images is not None
            pics = [c for c in slide.components if c.type == ComponentType.PICTURE]
            if pics:
                assert pics[0].id in images
                assert images[pics[0].id].startswith("data:image/")
            return {
                c.id: {"context": f"ctx for {c.id}", "sources": []}
                for c in slide.components
            }

    with (
        patch("app.services.analysis.EmbeddingsClient", FakeEmbed),
        patch("app.services.analysis.LLMClient", FakeLLM),
    ):
        deck = analyze_deck(
            slides_bytes=_pptx_with_picture(),
            slides_filename="infographic.pptx",
            doc_bytes=_docx_bytes(),
            doc_filename="notes.docx",
            settings=settings,
        )

    pics = [c for s in deck.slides for c in s.components if c.type == ComponentType.PICTURE]
    assert len(pics) == 1
    assert pics[0].has_image is True
    assert pics[0].id in deck.component_images
    assert deck.component_images[pics[0].id].startswith("data:image/")
    assert pics[0].context.startswith("ctx for")


def test_chat_passes_component_image_to_llm() -> None:
    """Session chat about a picture component includes image_data_uri."""
    session_store.clear()
    deck_store.clear()

    pic_id = "slide0-shape2"
    data_uri = "data:image/png;base64,qqq"
    deck = Deck(
        id="deck-img",
        slides_filename="a.pptx",
        doc_filename="a.docx",
        doc_filenames=["a.docx"],
        slides=[
            Slide(
                index=0,
                components=[
                    Component(
                        id=pic_id,
                        type=ComponentType.PICTURE,
                        bbox=BBox(left=0.2, top=0.2, width=0.5, height=0.5),
                        text="Infographic",
                        has_image=True,
                        context="A growth chart",
                    )
                ],
            )
        ],
        doc_chunks=[
            DocChunk(text="Enterprise drove growth.", source="p:0", embedding=[1.0, 0.0])
        ],
        component_images={pic_id: data_uri},
    )
    deck_store.put(deck)

    session = get_or_create_session(
        deck, CreateSessionRequest(component_id=pic_id, audience="general")
    )
    settings = Settings(mistral_api_key="test-key")
    captured: dict = {}

    def fake_answer_chat(**kwargs):
        captured.update(kwargs)
        return "The image shows enterprise growth."

    class FakeEmbed:
        def __init__(self, _s):
            pass

        def embed(self, texts):
            return [[1.0, 0.0] for _ in texts]

    with (
        patch("app.services.chat.EmbeddingsClient", FakeEmbed),
        patch("app.services.chat.LLMClient") as mock_llm_cls,
    ):
        mock_llm_cls.return_value.answer_chat.side_effect = fake_answer_chat
        # Need side_effect on instance method - simpler to patch answer_chat via return_value
        mock_llm_cls.return_value.answer_chat = MagicMock(
            side_effect=lambda **kw: captured.update(kw) or "ok"
        )
        send_message(
            deck,
            session,
            SendMessageRequest(mode=QueryMode.ASK, content="What does this show?"),
            settings,
        )

    assert captured.get("image_data_uri") == data_uri


def test_query_passes_component_image_to_llm() -> None:
    pic_id = "slide0-shape2"
    data_uri = "data:image/png;base64,zzz"
    deck = Deck(
        id="deck-q",
        slides_filename="a.pptx",
        doc_filename="a.docx",
        slides=[
            Slide(
                index=0,
                components=[
                    Component(
                        id=pic_id,
                        type=ComponentType.PICTURE,
                        bbox=BBox(left=0.1, top=0.1, width=0.3, height=0.3),
                        text="Pic",
                        has_image=True,
                    )
                ],
            )
        ],
        doc_chunks=[DocChunk(text="doc", source="p:0", embedding=[0.0, 1.0])],
        component_images={pic_id: data_uri},
    )
    settings = Settings(mistral_api_key="test-key")
    captured: dict = {}

    class FakeEmbed:
        def __init__(self, _s):
            pass

        def embed(self, texts):
            return [[0.0, 1.0] for _ in texts]

    with (
        patch("app.services.qa.EmbeddingsClient", FakeEmbed),
        patch("app.services.qa.LLMClient") as mock_cls,
    ):
        mock_cls.return_value.answer_query = MagicMock(
            side_effect=lambda **kw: captured.update(kw) or "answer"
        )
        answer_query(
            deck,
            QueryRequest(mode=QueryMode.ASK, question="Describe it", component_id=pic_id),
            settings,
        )

    assert captured.get("image_data_uri") == data_uri
