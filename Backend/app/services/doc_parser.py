from __future__ import annotations

from io import BytesIO
from typing import BinaryIO

from docx import Document
from pypdf import PdfReader

from app.models.schemas import DocChunk


def _extract_docx(file_bytes: bytes | BinaryIO) -> list[tuple[str, str]]:
    """Return list of (text, source_label) from a DOCX."""
    stream = BytesIO(file_bytes) if isinstance(file_bytes, (bytes, bytearray)) else file_bytes
    doc = Document(stream)
    passages: list[tuple[str, str]] = []

    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if text:
            passages.append((text, f"paragraph:{i}"))

    for t_idx, table in enumerate(doc.tables):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            passages.append(("\n".join(rows), f"table:{t_idx}"))

    return passages


def _extract_pdf(file_bytes: bytes | BinaryIO) -> list[tuple[str, str]]:
    """Return list of (text, source_label) from a PDF."""
    stream = BytesIO(file_bytes) if isinstance(file_bytes, (bytes, bytearray)) else file_bytes
    reader = PdfReader(stream)
    passages: list[tuple[str, str]] = []

    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            passages.append((text, f"page:{i}"))

    return passages


def extract_doc_passages(file_bytes: bytes, filename: str) -> list[tuple[str, str]]:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    if lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    raise ValueError(f"Unsupported document format: {filename}")


def chunk_text(
    passages: list[tuple[str, str]],
    max_chars: int = 1000,
    overlap: int = 150,
) -> list[DocChunk]:
    """
    Chunk passages into DocChunks of roughly max_chars, with overlap.
    Each chunk retains the source label of the originating passage.
    """
    chunks: list[DocChunk] = []

    for text, source in passages:
        if len(text) <= max_chars:
            chunks.append(DocChunk(text=text, source=source))
            continue

        start = 0
        part = 0
        while start < len(text):
            end = start + max_chars
            piece = text[start:end].strip()
            if piece:
                label = source if part == 0 else f"{source}#{part}"
                chunks.append(DocChunk(text=piece, source=label))
            if end >= len(text):
                break
            start = max(0, end - overlap)
            part += 1

    return chunks


def parse_document(
    file_bytes: bytes,
    filename: str,
    max_chars: int = 1000,
    overlap: int = 150,
) -> list[DocChunk]:
    passages = extract_doc_passages(file_bytes, filename)
    return chunk_text(passages, max_chars=max_chars, overlap=overlap)
